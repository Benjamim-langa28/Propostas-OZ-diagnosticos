import csv
import io
import unicodedata
from io import BytesIO

import fitz
from docx import Document
from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook

MAX_UPLOAD_BYTES = 15 * 1024 * 1024


def extract_mqt_rows(filename: str, data: bytes) -> list[dict[str, str | float | None]]:
    """Read common MQT spreadsheets where a header row has a description and quantity column."""
    lower_name = filename.lower()
    if lower_name.endswith(".xlsx"):
        book = load_workbook(BytesIO(data), read_only=True, data_only=True)
        sheets = [[list(row) for row in sheet.iter_rows(values_only=True)] for sheet in book.worksheets]
    elif lower_name.endswith(".csv"):
        decoded = data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(decoded[:4096], delimiters=";,\t")
        except csv.Error:
            dialect = csv.excel
        sheets = [[row for row in csv.reader(io.StringIO(decoded), dialect=dialect)]]
    else:
        return []

    def normalize(value: object) -> str:
        raw = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().lower()
        return " ".join(raw.replace(".", " ").replace("_", " ").split())

    def number(value: object) -> float | None:
        if value is None:
            return None
        try:
            if isinstance(value, (int, float)):
                return float(value)
            raw = str(value).strip().replace(" ", "")
            if "," in raw and "." in raw:
                raw = raw.replace(".", "").replace(",", ".")
            else:
                raw = raw.replace(",", ".")
            return float(raw)
        except ValueError:
            return None

    items: list[dict[str, str | float | None]] = []
    for rows in sheets:
        header_index = -1
        description_index = quantity_index = unit_index = code_index = None
        for index, row in enumerate(rows[:12]):
            labels = [normalize(value) for value in row]
            description_index = next((i for i, label in enumerate(labels) if any(word in label for word in
                ("descricao", "designacao", "description", "trabalho", "servico"))), None)
            quantity_index = next((i for i, label in enumerate(labels) if any(word in label for word in
                ("quantidade", "qtd", "quantity", "qty"))), None)
            if description_index is not None and quantity_index is not None:
                header_index = index
                unit_index = next((i for i, label in enumerate(labels) if any(word in label for word in
                    ("unidade", "unit"))), None)
                code_index = next((i for i, label in enumerate(labels) if any(word in label for word in
                    ("codigo", "referencia", "reference", "ref"))), None)
                break
        if header_index < 0 or description_index is None or quantity_index is None:
            continue
        for row in rows[header_index + 1:]:
            description = str(row[description_index]).strip() if description_index < len(row) and row[description_index] is not None else ""
            quantity = number(row[quantity_index]) if quantity_index < len(row) else None
            if not description or quantity is None or quantity < 0:
                continue
            code = str(row[code_index]).strip() if code_index is not None and code_index < len(row) and row[code_index] is not None else None
            unit = str(row[unit_index]).strip() if unit_index is not None and unit_index < len(row) and row[unit_index] is not None else "un"
            items.append({"code": code or None, "description": description[:500], "unit": unit[:50] or "un", "quantity": quantity})
    return items[:500]


async def extract_upload(file: UploadFile) -> str:
    data = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "O anexo excede o limite de 15 MB.")
    name = (file.filename or "anexo").lower()
    if name.endswith((".txt", ".csv", ".eml")):
        return data.decode("utf-8", errors="replace")
    if name.endswith(".pdf") or file.content_type == "application/pdf":
        try:
            pdf = fitz.open(stream=data, filetype="pdf")
            text = "\n\n".join(page.get_text("text") for page in pdf).strip()
        except Exception as exc:
            raise HTTPException(422, "Não foi possível ler este PDF.") from exc
        if not text:
            raise HTTPException(422, "Este PDF parece digitalizado. OCR ainda não está configurado.")
        return text
    if name.endswith(".docx"):
        try:
            doc = Document(BytesIO(data))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception as exc:
            raise HTTPException(422, "Não foi possível ler este documento Word.") from exc
    if name.endswith(".xlsx"):
        try:
            book = load_workbook(BytesIO(data), read_only=True, data_only=True)
            rows = []
            for sheet in book.worksheets:
                rows.append(f"Folha: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                    if values:
                        rows.append(" | ".join(values))
            return "\n".join(rows)
        except Exception as exc:
            raise HTTPException(422, "Não foi possível ler esta folha Excel.") from exc
    raise HTTPException(415, "Formato não suportado. Usa PDF, DOCX, XLSX, TXT, CSV ou EML.")
