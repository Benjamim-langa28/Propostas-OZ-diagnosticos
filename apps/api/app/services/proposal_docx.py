"""Gera a proposta em Word (.docx) no formato das propostas OZ."""
from __future__ import annotations

import io
import os
import re
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image as DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.config import settings

FONT = "Arial"
LEFT, CENTER, RIGHT, JUSTIFY = (WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
                                WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.JUSTIFY)

# Textos-padrão (revê com a equipa). Estão aqui para serem fáceis de alterar.
STANDARD_PRICE_NOTES = (
    "O preço proposto tem em conta a especificidade do tipo de trabalhos, executados por pessoal qualificado, "
    "sob a coordenação de um engenheiro civil com experiência em diagnóstico e reabilitação de estruturas, "
    "utilizando meios e equipamentos adequados para o efeito.",
    "Não se inclui a execução de ensaios adicionais, para além dos propostos, nem de outros trabalhos "
    "eventualmente necessários.",
    "As verbas indicadas não incluem o IVA à taxa legal em vigor.",
)
CONDITIONAL_NOTE = ("As atividades identificadas como condicionais executam-se se a condição técnica indicada se "
                    "verificar, sem nova aprovação do Cliente, e estão incluídas no valor global proposto.")
STANDARD_EXCLUSIONS = (
    "Autorizações necessárias.",
    "O pleno e livre acesso aos locais a inspecionar, removendo todos os obstáculos que impeçam ou dalguma forma o condicionem.",
    "Energia elétrica, monofásica, 220 V e água à pressão da rede.",
    "Reparações que não as indicadas.",
    "Tudo o que não estiver incluído na presente proposta.",
)
SUBSTITUTION_NOTE = (
    "Alguns dos trabalhos acima indicados poderão ser substituídos por outros equivalentes, caso se verifique, "
    "face às condições que virem a ser constatadas no local, ser desaconselhável a sua realização nos moldes "
    "preconizados ou haver vantagem no recurso a outras técnicas. Tal substituição será sempre submetida "
    "antecipadamente à aprovação do Cliente.")


def _cfg(name: str, default: str = "") -> str:
    value = getattr(settings, name.lower(), None)
    if value in (None, ""):
        value = os.environ.get(name.upper())
    text = str(value).strip().strip('"').strip() if value not in (None, "") else ""
    return text or default


def _num(value: Any, default: float = 0.0) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _eur(value: float) -> str:
    value = float(value)
    text = f"{value:,.0f}" if abs(value - round(value)) < 0.005 else f"{value:,.2f}"
    return text.replace(",", "\u00a0").replace(".", ",") + " €"


def _qty(value: Any) -> str:
    number = _num(value)
    return str(int(number)) if number == int(number) else f"{number:g}".replace(".", ",")


def _lines(value: Any) -> list[str]:
    if isinstance(value, (list, tuple)):
        return [str(v).strip() for v in value if str(v).strip()]
    return [p.strip() for p in str(value or "").splitlines() if p.strip()]


def _sentences(text: str) -> list[str]:
    return [s.strip() for s in re.split(r"(?<=[.!?])\s+", text or "") if s.strip()]


def _condition_sentence(condition: Any) -> str:
    text = " ".join(str(condition or "").split()).rstrip(".")
    if not text:
        return ""
    if re.search(r"\b(PAT|ENS|CAU|SOL)-[A-Z0-9-]+", text):
        return "Atividade condicional, a realizar se a inspeção de campo confirmar indícios que a justifiquem."
    if text[0].islower():
        return f"Atividade condicional, a realizar se {text}."
    return f"Atividade condicional. {text}."


def _header_image() -> Path | None:
    folder = Path(__file__).resolve().parent.parent / "assets"
    if not folder.is_dir():
        return None
    images = [p for p in folder.rglob("*") if p.suffix.lower() in {".png", ".jpg", ".jpeg"}
              and not any(k in p.name.lower() for k in ("assin", "sign"))]  # nunca reutilizar a assinatura antiga
    if not images:
        return None
    images.sort(key=lambda p: (0 if any(k in p.name.lower() for k in ("cabecalho", "header", "logo", "oz")) else 1, p.name))
    return images[0]


def _add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = instruction
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, end):
        run._r.append(element)


def _shade(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear"); shading.set(qn("w:color"), "auto"); shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _cell(cell, text: Any = "", *, bold=False, italic=False, size=9.5, align=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(str(text))
    run.bold, run.italic, run.font.size = bold, italic, Pt(size)
    return paragraph


def _label_cell(cell, pt: str, en: str = "") -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    first = paragraph.add_run(pt); first.bold = True; first.font.size = Pt(9.5)
    if en:
        second = paragraph.add_run(f" ({en})"); second.bold = True; second.italic = True; second.font.size = Pt(9.5)
    last = paragraph.add_run(":"); last.bold = True; last.font.size = Pt(9.5)


def _para(document, text="", *, bold=False, italic=False, size=None, align=JUSTIFY, after=None,
          color=None, indent=None, keep_next=False):
    paragraph = document.add_paragraph()
    if text:
        run = paragraph.add_run(text)
        run.bold, run.italic = bold, italic
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor(*color)
    paragraph.alignment = align
    if after is not None:
        paragraph.paragraph_format.space_after = Pt(after)
    if indent is not None:
        paragraph.paragraph_format.left_indent = Cm(indent)
    if keep_next:
        paragraph.paragraph_format.keep_with_next = True
    return paragraph


def _heading(document, number: str, title: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = LEFT
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run(f"{number}.").bold = True
    paragraph.add_run(f"\u2003{title}")


def _dash(document, text: str) -> None:
    _para(document, f"–\u2002{text}", indent=0.6, after=3)


def _setup_page(document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.5)
    section.top_margin = section.bottom_margin = Cm(2.0)
    section.header_distance, section.footer_distance = Cm(1.0), Cm(0.8)

    header_paragraph = section.header.paragraphs[0]
    image = _header_image()
    done = False
    if image:
        try:
            picture = DocxImage.from_file(str(image))
            if picture.px_height / picture.px_width <= 0.35:
                header_paragraph.add_run().add_picture(str(image), width=Cm(16.0))
            else:
                header_paragraph.add_run().add_picture(str(image), height=Cm(2.6))
            done = True
        except Exception:
            done = False
    if not done:
        header_paragraph.alignment = RIGHT
        run = header_paragraph.add_run(_cfg("COMPANY_NAME", "Diagnóstico, Levantamento e Controlo de Qualidade em Estruturas e Fundações, Lda."))
        run.bold, run.font.size = True, Pt(9)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    run = footer_paragraph.add_run(
        f"Escritório: {_cfg('COMPANY_ADDRESS', 'Rua Prof. Reinaldo dos Santos, 48 - B, 1500-508 Lisboa')} • "
        f"Tel.: {_cfg('COMPANY_PHONE', '213 563 371')}  Fax: {_cfg('COMPANY_FAX', '213 153 550')}\n"
        f"Email: {_cfg('COMPANY_EMAIL', 'ger@oz-diagnostico.pt')} • URL: {_cfg('COMPANY_WEBSITE', 'www.oz-diagnostico.pt')}")
    run.font.size = Pt(7.5)
    page_paragraph = footer.add_paragraph()
    page_paragraph.alignment = RIGHT
    _add_field(page_paragraph, "PAGE")
    page_paragraph.add_run("/")
    _add_field(page_paragraph, "NUMPAGES")


def _reference_table(document, fields, proposal_no) -> None:
    table = document.add_table(rows=0, cols=4)
    table.autofit = False
    widths = [3.6, 5.2, 2.6, 4.6]
    signatory = _cfg("PROPOSAL_SIGNATORY") or _cfg("COMPANY_NAME", "")
    rows = [
        (("N. ref.", "Our ref."), proposal_no, ("Data", "Date"), date.today().isoformat()),
        (("Para", "To"), fields.get("client_name") or "", ("Tel.", "Phone"), fields.get("client_phone") or ""),
        (("À At.", "ATT."), fields.get("client_contact_name") or "", ("Email", ""),
         fields.get("client_contact_email") or fields.get("client_email") or ""),
        (("CC.", ""), "", None, ""),
        (("DE", "From"), signatory, ("Tel.", "Phone"), _cfg("COMPANY_PHONE", "213 563 371")),
    ]
    for (l1, v1, l2, v2) in rows:
        cells = table.add_row().cells
        for cell, width in zip(cells, widths):
            cell.width = Cm(width)
        _label_cell(cells[0], *l1)
        _cell(cells[1], v1)
        if l2:
            _label_cell(cells[2], *l2)
        _cell(cells[3], v2)
    cells = table.add_row().cells
    for cell, width in zip(cells, widths):
        cell.width = Cm(width)
    _label_cell(cells[0], "ASS.", "RE")
    subject = fields.get("source_subject") or fields.get("project_name") or ""
    objective = " ".join(str(fields.get("objective") or "").split()).rstrip(".")
    merged = cells[1].merge(cells[3])
    merged.text = ""
    first = merged.paragraphs[0]
    first.paragraph_format.space_after = Pt(2)
    first.add_run(f"V. solicitação: “{subject}”.").font.size = Pt(9.5)
    if objective:
        second = merged.add_paragraph()
        second.paragraph_format.space_after = Pt(2)
        second.add_run(f"Proposta de prestação de serviços de {objective[0].lower() + objective[1:]}.").font.size = Pt(9.5)


def _annex(document, active, vat_rate, total, fields) -> None:
    document.add_page_break()
    _para(document, "MAPA DE QUANTIDADES E PREÇOS UNITÁRIOS", bold=True, align=CENTER, after=2)
    subtitle = fields.get("location") or fields.get("project_name")
    if subtitle:
        _para(document, str(subtitle), italic=True, align=CENTER, after=8)
    widths = [1.1, 7.2, 1.4, 1.3, 2.5, 2.5]
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, (cell, title) in enumerate(zip(table.rows[0].cells,
                                              ["Item", "Designação", "Un.", "Qt.", "Preço Unitário", "Preço Total"])):
        cell.width = Cm(widths[index])
        _cell(cell, title, bold=True, align=LEFT if index == 1 else CENTER)
        _shade(cell, "D9D9D9")
    for number, item in enumerate(active, start=1):
        cells = table.add_row().cells
        quantity, price = _num(item.get("quantity")), _num(item.get("unit_price"))
        pending = price <= 0
        name = str(item.get("name") or "")
        if item.get("activity_class") == "CONDITIONAL":
            name += " (condicional)"
        values = [f"{number}.", name, str(item.get("unit") or ""), _qty(quantity),
                  "A validar" if pending else _eur(price), "—" if pending else _eur(quantity * price)]
        for index, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Cm(widths[index])
            _cell(cell, value, align=RIGHT if index >= 4 else (LEFT if index == 1 else CENTER))
            if pending and index >= 4:
                _shade(cell, "FFF2CC")
    vat = total * vat_rate / 100
    for label, amount, strong in (("TOTAL", total, True), (f"IVA ({_qty(vat_rate)}%)", vat, False),
                                  ("TOTAL COM IVA", total + vat, True)):
        cells = table.add_row().cells
        for cell, width in zip(cells, widths):
            cell.width = Cm(width)
        merged = cells[0].merge(cells[4])
        _cell(merged, label, bold=strong, align=RIGHT)
        _cell(cells[5], _eur(amount), bold=strong, align=RIGHT)

def _rank(item: dict[str, Any]) -> int:
    """Ordem lógica: mobilização, inspeção, ensaios, trabalhos intrusivos, relatório no fim."""
    sid = str(item.get("service_id") or "")
    name = str(item.get("name") or "").lower()
    if sid in {"MOB", "VERBA_LEV"}:
        return 0
    if sid in {"EST_VISUAL", "INSP_ANOM", "LEV_ANOM", "MAP_FISS", "LEV_DIM"}:
        return 1
    if sid.startswith("REL") or "relatório" in name:
        return 9
    if sid in {"ABERT_LOCAL", "EXPOS_ARM", "REM_REVEST", "INSP_INT", "FECHO_ABERT", "REP_REVEST"}:
        return 5
    return 3

def build_proposal_docx(request_row: dict[str, Any], proposal: dict[str, Any],
                        items: list[dict[str, Any]], analysis: dict[str, Any] | None = None) -> bytes:
    fields = request_row.get("extracted_fields") or {}
    active = [i for i in items if i.get("enabled", True)]
    active.sort(key=lambda r: (_rank(r), _num(r.get("position")), str(r.get("created_at") or "")))
    vat_rate = _num(proposal.get("vat_rate"), 23.0)
    pending = [i for i in active if _num(i.get("unit_price")) <= 0]
    total = sum(_num(i.get("quantity")) * _num(i.get("unit_price")) for i in active if _num(i.get("unit_price")) > 0)
    proposal_no = str(proposal.get("proposal_no") or "")

    document = Document()
    document.core_properties.title = f"Proposta {proposal_no}"
    document.core_properties.author = _cfg("COMPANY_NAME", "OZ")
    _setup_page(document)

    if pending:
        _para(document, f"RASCUNHO PARA REVISÃO: {len(pending)} serviço(s) sem preço validado (assinalados a amarelo no "
                        "mapa de preços). Remover este aviso antes de enviar.", bold=True, color=(192, 0, 0), size=9, align=LEFT)
    _reference_table(document, fields, proposal_no)
    _para(document, f"Proposta Nº.: {proposal_no}", bold=True, align=LEFT, after=10).paragraph_format.space_before = Pt(10)
    _para(document, "Exºs. Senhores,", align=LEFT)
    _para(document, "Na sequência da V. solicitação e posteriores esclarecimentos, sobre o assunto em apreço, que agradecemos, "
                    "apresentamos as condições em que poderemos vir a prestar a nossa colaboração.")

    _heading(document, "1", "Objeto")
    _para(document, str(fields.get("objective") or request_row.get("title") or "A definir."))
    project = " ".join(str(fields.get("project_name") or "").split())
    location = " ".join(str(fields.get("location") or "").split())
    if re.match(r"(?i)^(para\s+)?(inspe|diagn|pedido|proposta|ensaio)", project) or len(project) > 90:
        project = ""  # parece o assunto do email e não o nome da obra
    place = ", ".join(p for p in (project, location) if p)
    if place:
        _para(document, f"Obra: {place}.")

    _heading(document, "2", "Considerações prévias")
    if fields.get("problem_summary"):
        _para(document, str(fields["problem_summary"]))
    requested = _lines(fields.get("requested_services"))
    if requested:
        _para(document, "De acordo com o pedido recebido, pretende-se:", after=3)
        for line in requested:
            _dash(document, line)
    conditions = _lines(fields.get("requested_conditions"))
    if conditions:
        _para(document, "Pontos adicionais solicitados pelo Cliente, a esclarecer antes do envio:", after=3)
        for line in conditions:
            _dash(document, line)
    _para(document, "Descrevem-se, a seguir, as condições técnicas da prestação de serviços.")

    _heading(document, "3", "Condições técnicas")
    if not active:
        _para(document, "Âmbito por definir.")
    for number, item in enumerate(active, start=1):
        paragraph = document.add_paragraph()
        paragraph.alignment = LEFT
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run(f"3.{number}.\u2002{item.get('name') or ''}").bold = True
        if item.get("technical_basis"):
            basis = re.split(r"\s*Condição de execução:", str(item["technical_basis"]))[0].strip()
            if basis and "Aderência textual" not in basis and "base técnica" not in basis:
                _para(document, basis, after=3)
        _para(document, f"Quantidade prevista: {_qty(item.get('quantity'))} {item.get('unit') or ''}.".strip(), after=3)
        if item.get("activity_class") == "CONDITIONAL" and item.get("trigger_condition"):
            _para(document, _condition_sentence(item["trigger_condition"]), italic=True, after=3)
    _para(document, SUBSTITUTION_NOTE)

    _heading(document, "4", "Preços")
    _para(document, f"Para a prestação da nossa colaboração nos moldes anteriormente descritos, propõe-se a verba global de "
                    f"{_eur(total)} (acresce IVA à taxa legal em vigor), apresentando-se em anexo o mapa de quantidades e preços unitários.")
    _para(document, "Notas:", after=3)
    notes = list(STANDARD_PRICE_NOTES)
    if any(i.get("activity_class") == "CONDITIONAL" for i in active):
        notes.insert(0, CONDITIONAL_NOTE)
    for note in notes:
        _dash(document, note)

    _heading(document, "5", "Condições de pagamento")
    payment = _cfg("PROPOSAL_PAYMENT_TERMS")
    for sentence in (_sentences(payment) or ["[Condições de pagamento a definir]"]):
        _dash(document, sentence)

    _heading(document, "6", "Prazos")
    execution = str(proposal.get("execution_period") or "").strip() or _cfg("PROPOSAL_EXECUTION_PERIOD")
    for sentence in (_sentences(execution) or ["Início dos trabalhos: a combinar."]):
        _para(document, sentence, after=3)
    _para(document, f"Validade da proposta: {int(_num(proposal.get('validity_days'), 60))} dias.")

    _heading(document, "7", "Exclusões")
    _para(document, "São de conta de vossas Exªs, as seguintes tarefas e fornecimentos (ver também, ponto 4):", after=3)
    for line in STANDARD_EXCLUSIONS:
        _dash(document, line)

    quality = _cfg("PROPOSAL_QUALITY_STATEMENT")
    affiliations = _cfg("PROPOSAL_AFFILIATIONS_STATEMENT")
    if quality or affiliations:
        _para(document, " ".join(p for p in (quality, affiliations) if p)).paragraph_format.space_before = Pt(10)
    _para(document, "Esperando que a nossa proposta vá ao encontro dos vossos interesses, aproveitamos para apresentar os "
                    "nossos melhores cumprimentos.")
    _para(document, "Atentamente,", align=LEFT, after=36)
    _para(document, f"({_cfg('PROPOSAL_SIGNATORY') or 'Assinatura a configurar'})", align=LEFT)

    _annex(document, active, vat_rate, total, fields)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()